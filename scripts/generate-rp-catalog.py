#!/usr/bin/env python3
"""Genere rp-catalog.ts (bot + API) et rp-grades.ts (site) depuis Google Sheets / Excel Site-12."""

from __future__ import annotations

import re
import sys
import unicodedata
import urllib.request
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
SHEETS_ID = "1xHyhay1CWD3J03KQzrXeu2tZH8rupdEu3ndxzK0J0EE"
EXCEL_DEFAULT = Path.home() / "Downloads" / "Branche Du Site 12 (4).xlsx"
OUT_BOT = ROOT / "apps/discord-bot/src/lib/rp-catalog.ts"
OUT_API = ROOT / "apps/api/src/sync/rp-catalog.ts"
OUT_WEB = ROOT / "apps/web/src/data/rp-grades.ts"
OUT_DISCORD_GEN = ROOT / "apps/discord-bot/src/lib/discord-master-catalog.generated.ts"
OUT_DISCORD_MD = ROOT / "docs/DISCORD-GRADES-COMPLET.md"

SKIP_RE = re.compile(
    r"^(TEAM |CHAMBRE |EXPERIENCE |OQUIPE |O-\d|EQUIPE |Objectif|Utiliter|"
    r"PAYE|Paye|GRADES|ACCES|Arm\d|A\d$|N\d$|W\.H\.|O5$|Keter|Euclid|Safe|"
    r"SCP$|Soin$|Archives$|Qualifier$|Agu|Compétent$|Experimenter$|Novice$|"
    r"Recherche$|Entretiens$|Réparation$|Livraison$|Secretairiat$|"
    r"Restauration$|Nettoyage$|Communication$|Commande$|Construction$|"
    r"Sécuriter$|Class-D / B / S$|Arme |Carte |Mun$|Melee$|Armements$|"
    r"Agir$|Avertir$|Controler$|Conseiller$|Contenir$|Décider$|Diriger$|"
    r"Employé$|Etudier$|Eviter$|Garantir$|Financer$|Soutenir$|Surperviser$|"
    r"Check$|Gates$|Inter\.$|Maint\.$|Technologie$|Autorisation|Inspecteur$|"
    r"Salaires|Salaires Total|Total Salaires)",
    re.I,
)

BRANCH_MAP = {
    "omega": [
        "gerant omega",
        "adjoint omega",
        r"omega-o\d",
        r"omega-0\d",
        r"president du conseil",
        r"vice-president",
        r"conseiller",
    ],
    "direction": [
        "directeur du site",
        "adjoint-directeur$",
        "adjoint directeur$",
    ],
    "securite": [
        "directeur secur", "adjoint-directeur se", "commandant", "lieutenant",
        "superviseur", "sergent", "caporal", "soldat", "garde", "general de com",
        "instructeur", "inspecteur", "responsable garde", "recrue",
    ],
    "scientifique": [
        "directeur scient", "adjoint-directeur sc", "scientifique", "chercheur",
        "archiviste", "dr.", "medecin", "psychologue", "m/c dr", "responsable d'autorisation",
        "chef d'archive", "chef de com", "superviseur d'experience", "superviseur de scp",
        "superviseur d'intervention", "directeur-adjoint medical", "medecin en chef",
    ],
    "maintenance": [
        "directeur maint", "adjoint-directeur maint", "responsable d'entretien",
        "responsable livraison", "chef d'intervention", "chef de commande",
        "com. de maintenance", "chef technicien", "technicien", "plombier",
        "electricien", "mecanicien", "chef receptionniste", "demenageur", "travailleur",
    ],
    "general": [
        "directeur general", "adjoint-directeur-general", "secretaire", "com. ",
        "superviseur secretaire", "superviseur restauration", "superviseur nettoyage",
        "superviseur communication", "chef de cuisine", "concierge",
    ],
    "classes": ["class -", "class-"],
}

TIER_ACCESS = {
    "omega": ["wh", "o5", "gates", "n5", "keter", "a5", "a4", "euclid", "a3", "a2", "safe", "a1", "n4", "n3", "n2", "n1", "arm3", "arm2", "arm1", "check", "inter", "maint"],
    "direction": ["wh", "o5", "gates", "n5", "keter", "a5", "a4", "euclid", "a3", "a2", "safe", "a1", "n4", "n3", "n2", "n1", "check", "inter", "maint"],
    "officier": ["gates", "n5", "keter", "a5", "a4", "euclid", "a3", "a2", "safe", "a1", "n4", "n3", "n2", "check", "inter"],
    "sous-officier": ["n4", "n3", "n2", "n1", "a2", "a1", "safe", "euclid", "check"],
    "troupe": ["n2", "n1", "a1", "safe", "check"],
    "scientifique": ["n5", "keter", "euclid", "safe", "n4", "n3", "n2", "maint"],
    "medical": ["n4", "n3", "n2", "inter", "maint"],
    "technique": ["n3", "n2", "n1", "maint"],
    "admin": ["n3", "n2", "n1", "check"],
    "class": ["n1"],
}

# Conseil Oméga — 5 sièges (O1→O5), noms affichés cohérents
OMEGA_SKIP_KEYS = frozenset({
    "gerant omega-01",
    "adjoint omega-02",
    "omega-03 a 05",
})

OMEGA_CANONICAL_NAMES: dict[str, str] = {
    "gerant omega-o1": "Président du Conseil (O1)",
    "adjoint omega-o2": "Vice-Président — Sécurité (O2)",
    "omega-o3": "Conseiller — Sciences (O3)",
    "omega-o4": "Conseiller — Maintenance (O4)",
    "omega-o5": "Conseiller — Services (O5)",
}


DIRECTION_SKIP_KEYS = frozenset({"directeur branche"})


def apply_omega_entry(name: str, key: str) -> tuple[str, str] | None:
    """Retourne (nom affiché, clé) ou None si entrée à ignorer."""
    if key in OMEGA_SKIP_KEYS:
        return None
    if key in OMEGA_CANONICAL_NAMES:
        name = OMEGA_CANONICAL_NAMES[key]
        return name, normalize(name)
    return name, key


def finalize_omega_grade(entry: dict) -> dict:
    key = normalize(entry["name"])
    if key not in {normalize(v) for v in OMEGA_CANONICAL_NAMES.values()}:
        return entry
    entry = {
        **entry,
        "branch": "omega",
        "tier": "omega",
        "departmentId": "omega",
        "clearance": clearance_for("omega"),
        "accessZones": TIER_ACCESS["omega"],
        "siteSections": site_sections("omega", "omega"),
    }
    entry["description"] = build_description(entry)
    return entry


def normalize(value: str) -> str:
    return (
        unicodedata.normalize("NFD", value)
        .encode("ascii", "ignore")
        .decode("ascii")
        .lower()
        .replace(" - ", "-")
        .replace(" -", "-")
        .replace("- ", "-")
        .replace("  ", " ")
        .strip()
    )


def safe_int(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    s = str(value).strip()
    if not s or not re.match(r"^[\d.]+$", s):
        return None
    return int(float(s))


def download_sheets(path: Path) -> None:
    url = f"https://docs.google.com/spreadsheets/d/{SHEETS_ID}/export?format=xlsx"
    print(f"Telechargement Google Sheets -> {path}")
    urllib.request.urlretrieve(url, path)


def resolve_excel_path(argv: list[str]) -> Path:
    if len(argv) > 1:
        return Path(argv[1])
    cache = ROOT / "scripts/.cache/site12.xlsx"
    cache.parent.mkdir(parents=True, exist_ok=True)
    if not cache.exists() and not EXCEL_DEFAULT.exists():
        download_sheets(cache)
        return cache
    if EXCEL_DEFAULT.exists():
        return EXCEL_DEFAULT
    if not cache.exists():
        download_sheets(cache)
    return cache


def detect_branch(name: str) -> str:
    n = normalize(name)
    for branch, patterns in BRANCH_MAP.items():
        for p in patterns:
            if re.search(p, n):
                return branch
    return "general"


def detect_tier(name: str, branch: str) -> str:
    n = normalize(name)
    if branch == "omega":
        return "omega"
    if branch == "direction":
        return "direction"
    if branch == "classes":
        return "class"
    if branch == "scientifique":
        if re.search(r"medecin|psychologue|dr\.|m/c", n):
            return "medical"
        return "scientifique"
    if branch == "maintenance":
        return "technique"
    if branch == "general":
        if re.search(r"secretaire|com\.|superviseur", n):
            return "admin"
        return "admin"
    if re.search(r"commandant|lieutenant|superviseur|responsable|general", n):
        return "officier"
    if re.search(r"sergent|caporal", n):
        return "sous-officier"
    if re.search(r"soldat|recrue", n):
        return "troupe"
    return "troupe"


def site_sections(branch: str, tier: str) -> list[str]:
    sections = ["overview", "fondation"]
    if branch in ("omega", "direction"):
        sections += ["omega", "direction", "departements", "access-matrix"]
    if branch == "securite":
        sections += ["securite", "teams", "armory", "access-matrix"]
    if branch == "scientifique":
        sections += ["scientifique", "chambers", "experiences", "medical", "access-matrix"]
    if branch == "maintenance":
        sections += ["maintenance", "teams", "access-matrix"]
    if branch == "general":
        sections += ["general", "secretariat", "services"]
    if branch == "classes":
        sections += ["detention"]
    if tier in ("omega", "direction"):
        sections += ["mtf", "transmissions-classified"]
    elif tier in ("officier", "sous-officier"):
        sections += ["transmissions-restricted"]
    else:
        sections += ["transmissions-public"]
    return sorted(set(sections))


def clearance_for(tier: str) -> int:
    return {
        "omega": 5,
        "direction": 4,
        "officier": 3,
        "sous-officier": 2,
        "scientifique": 3,
        "medical": 2,
        "technique": 2,
        "admin": 2,
        "troupe": 1,
        "class": 1,
    }.get(tier, 1)


def build_description(entry: dict) -> str:
    parts = []
    branch = entry["branch"]
    dept_labels = {
        "omega": "Conseil Oméga",
        "direction": "Direction du Site-12",
        "securite": "Département Sécurité",
        "scientifique": "Département Scientifique",
        "maintenance": "Département Maintenance",
        "general": "Département Général",
        "classes": "Personnel détenu (Class-D/B/S)",
    }
    parts.append(f"Grade RP Site-12 — {dept_labels.get(branch, branch)}.")
    if entry.get("objectives"):
        parts.append("Missions : " + ", ".join(entry["objectives"][:6]) + ".")
    if entry.get("utilities"):
        parts.append("Domaines : " + ", ".join(entry["utilities"][:6]) + ".")
    if entry.get("pay"):
        parts.append(f"Rémunération hebdomadaire : {entry['pay']:,} $.".replace(",", " "))
    if entry.get("quota"):
        parts.append(f"Quota d'effectif : {entry['quota']}.")
    return " ".join(parts)


def parse_foundation_sheet(ws) -> dict[str, dict]:
    zones = [
        (c, str(ws.cell(5, c).value).strip())
        for c in range(28, 72)
        if ws.cell(5, c).value
    ]
    grades: dict[str, dict] = {}

    def ingest(name: str, pay, quota, row: int, objectives: list[str], utilities: list[str]):
        name = name.strip()
        if not name or SKIP_RE.search(name):
            return
        key = normalize(name)
        if not key:
            return
        if key in DIRECTION_SKIP_KEYS:
            return
        omega = apply_omega_entry(name, key)
        if omega is None:
            return
        name, key = omega
        access = {}
        for c, zone in zones:
            v = ws.cell(row, c).value
            if v and str(v).strip() not in ("", "Inspecteur", "ACCES"):
                access[zone] = str(v).strip()
        branch = detect_branch(name)
        tier = detect_tier(name, branch)
        computed_access = list(access.keys()) if access else TIER_ACCESS.get(tier, ["n1"])
        entry = {
            "name": name,
            "pay": safe_int(pay),
            "quota": safe_int(quota),
            "objectives": objectives,
            "utilities": utilities,
            "accessZones": computed_access,
            "branch": branch,
            "tier": tier,
            "clearance": clearance_for(tier),
            "siteSections": site_sections(branch, tier),
            "departmentId": {
                "securite": "securite",
                "scientifique": "recherche",
                "maintenance": "maintenance",
                "general": "general",
                "omega": "omega",
                "direction": "direction",
                "classes": "detention",
            }.get(branch, "general"),
        }
        entry["description"] = build_description(entry)
        prev = grades.get(key)
        if not prev or (entry["pay"] and not prev.get("pay")):
            grades[key] = entry

    for r in range(3, 260):
        name = ws.cell(r, 22).value
        if not name:
            continue
        ingest(
            str(name),
            ws.cell(r, 24).value,
            ws.cell(r, 26).value,
            r,
            [],
            [],
        )

    for r in range(3, 200):
        name = ws.cell(r, 4).value
        if not name:
            continue
        name = str(name).strip()
        if SKIP_RE.search(name):
            continue
        objectives: list[str] = []
        utilities: list[str] = []
        for c in range(6, 22, 2):
            v = ws.cell(r, c).value
            if not v:
                continue
            s = str(v).strip()
            if "Objectif" in s:
                continue
            if "Utiliter" in s:
                for c2 in range(c + 2, 22, 2):
                    u = ws.cell(r, c2).value
                    if u and "Objectif" not in str(u):
                        utilities.append(str(u).strip())
                break
            objectives.append(s)
        ingest(name, ws.cell(r, 24).value, ws.cell(r, 26).value, r, objectives, utilities)

    return grades


def from_excel(path: Path) -> tuple[list[str], dict[str, dict]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    grades: dict[str, dict] = {}
    if "Branche de la Fondation" in wb.sheetnames:
        grades.update(parse_foundation_sheet(wb["Branche de la Fondation"]))

    grade_cols = {4, 8, 9, 10, 11, 12}
    for sheet in ("Branche de la Fondation", "Description Tenue et Equipement"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            for idx, cell in enumerate(row, start=1):
                if idx not in grade_cols or cell is None:
                    continue
                s = str(cell).strip()
                if len(s) < 2 or SKIP_RE.search(s):
                    continue
                if re.match(r"^[\d.]+$", s):
                    continue
                key = normalize(s)
                if key in OMEGA_SKIP_KEYS or key in DIRECTION_SKIP_KEYS:
                    continue
                omega = apply_omega_entry(s, key)
                if omega is None:
                    continue
                s, key = omega
                if key not in grades:
                    branch = detect_branch(s)
                    tier = detect_tier(s, branch)
                    grades[key] = {
                        "name": s,
                        "pay": None,
                        "quota": None,
                        "objectives": [],
                        "utilities": [],
                        "accessZones": TIER_ACCESS.get(tier, ["n1"]),
                        "branch": branch,
                        "tier": tier,
                        "clearance": clearance_for(tier),
                        "siteSections": site_sections(branch, tier),
                        "departmentId": detect_branch(s),
                    }
                    grades[key]["description"] = build_description(grades[key])

    grades = {k: finalize_omega_grade(g) for k, g in grades.items()}
    names = sorted((g["name"] for g in grades.values()), key=normalize)
    return names, grades


def score_display_name(name: str) -> int:
    score = 0
    if re.search(r"[\u00c0-\u017f]", name):
        score += 12
    if re.search(r"\b(se\.|sc\.|maint\.)\b", name, re.I):
        score -= 8
    if name.endswith(".") and len(name) < 20:
        score -= 4
    if name == name.upper() and len(name) > 5:
        score -= 3
    return score


def dedupe_names(names: list[str]) -> list[str]:
    seen: set[str] = set()
    best: dict[str, str] = {}
    for n in names:
        key = normalize(n)
        if not key or key in seen and key not in best:
            continue
        prev = best.get(key)
        if not prev or score_display_name(n) > score_display_name(prev):
            best[key] = n
    return sorted(best.values(), key=normalize)


def render_catalog_ts(names: list[str], source: str) -> str:
    lines = [
        "/**",
        " * Grades RP Site-12 — synchronises sur Discord uniquement.",
        f" * Genere depuis : {source}",
        " * Staff, Membre, Joueur, Civil ne sont PAS geres par le bot.",
        " * Regenerer : python scripts/generate-rp-catalog.py",
        " */",
        "",
        "export const RP_GRADE_NAMES: readonly string[] = [",
    ]
    for n in names:
        esc = n.replace("\\", "\\\\").replace('"', '\\"')
        lines.append(f'  "{esc}",')
    lines.extend(
        [
            "];",
            "",
            "export const STAFF_ROLE_PATTERNS: readonly RegExp[] = [",
            '  /staff/i, /admin/i, /\\bmod\\b/i, /moderat/i, /owner/i, /fondateur/i,',
            '  /developpeur/i, /developer/i, /bot\\b/i, /verified|verifi/i,',
            '  /^membre$/i, /^joueur$/i, /^visiteur$/i, /^civil$/i,',
            '  /^nouveau/i, /^invite/i, /muted/i, /ping/i, /announcement/i,',
            '  /everyone/i, /technique\\s*staff/i, /gestion\\s*staff/i, /🛠/, /🛡.*staff/i,',
            "];",
            "",
            "export function normalizeRoleLabel(value: string): string {",
            '  return value.normalize("NFD").replace(/[\\u0300-\\u036f]/g, "")',
            '    .toLowerCase().replace(/\\s*-\\s*/g, "-")',
            '    .replace(/[^\\w\\s/.-]/g, "").replace(/\\s+/g, " ").trim();',
            "}",
            "",
            "export function isStaffOrBaseRole(roleName: string): boolean {",
            "  const n = normalizeRoleLabel(roleName);",
            "  if (!n) return true;",
            "  return STAFF_ROLE_PATTERNS.some((p) => p.test(roleName) || p.test(n));",
            "}",
            "",
            "export function isRpGradeName(name: string): boolean {",
            "  const n = normalizeRoleLabel(name);",
            "  return RP_GRADE_NAMES.some((g) => normalizeRoleLabel(g) === n);",
            "}",
            "",
            "function scoreDisplayName(name: string): number {",
            "  let score = 0;",
            '  if (/[\\u00c0-\\u017f]/.test(name)) score += 12;',
            '  if (/\\b(se\\.|sc\\.|maint\\.)\\b/i.test(name)) score -= 8;',
            "  if (name.endsWith(\".\") && name.length < 20) score -= 4;",
            "  if (name === name.toUpperCase() && name.length > 5) score -= 3;",
            "  return score;",
            "}",
            "",
            "export function getCanonicalGradeLabels(): Map<string, string> {",
            "  const map = new Map<string, string>();",
            "  for (const label of RP_GRADE_NAMES) {",
            "    const key = normalizeRoleLabel(label);",
            "    const prev = map.get(key);",
            "    if (!prev || scoreDisplayName(label) > scoreDisplayName(prev)) {",
            "      map.set(key, label);",
            "    }",
            "  }",
            "  return map;",
            "}",
            "",
        ]
    )
    return "\n".join(lines)


def render_web_ts(grades: dict[str, dict], source: str) -> str:
    entries = sorted(grades.values(), key=lambda g: normalize(g["name"]))
    lines = [
        "/**",
        f" * Metadonnees grades RP Site-12 — genere depuis {source}",
        " * Regenerer : python scripts/generate-rp-catalog.py",
        " */",
        "",
        "export type RpBranch =",
        '  | "omega" | "direction" | "securite" | "scientifique"',
        '  | "maintenance" | "general" | "classes";',
        "",
        "export type RpTier =",
        '  | "omega" | "direction" | "officier" | "sous-officier" | "troupe"',
        '  | "scientifique" | "medical" | "technique" | "admin" | "class";',
        "",
        "export interface RpGradeMeta {",
        "  id: string;",
        "  name: string;",
        "  branch: RpBranch;",
        "  tier: RpTier;",
        "  departmentId: string;",
        "  pay: number | null;",
        "  quota: number | null;",
        "  clearance: 1 | 2 | 3 | 4 | 5;",
        "  description: string;",
        "  objectives: string[];",
        "  utilities: string[];",
        "  accessZones: string[];",
        "  siteSections: string[];",
        "}",
        "",
        "export const rpGrades: RpGradeMeta[] = [",
    ]
    for g in entries:
        obj = {
            "id": normalize(g["name"]).replace(" ", "-"),
            "name": g["name"],
            "branch": g["branch"],
            "tier": g["tier"],
            "departmentId": g.get("departmentId", g["branch"]),
            "pay": g["pay"],
            "quota": g["quota"],
            "clearance": g["clearance"],
            "description": g["description"],
            "objectives": g["objectives"],
            "utilities": g["utilities"],
            "accessZones": g["accessZones"],
            "siteSections": g["siteSections"],
        }
        import json

        lines.append(f"  {json.dumps(obj, ensure_ascii=False)},")
    lines.extend(
        [
            "];",
            "",
            "export function normalizeGradeId(name: string): string {",
            '  return name.normalize("NFD").replace(/[\\u0300-\\u036f]/g, "")',
            '    .toLowerCase().replace(/\\s*-\\s*/g, "-")',
            '    .replace(/[^\\w\\s/.-]/g, "").replace(/\\s+/g, " ").trim();',
            "}",
            "",
            "export function findGradeMeta(name: string): RpGradeMeta | undefined {",
            "  const key = normalizeGradeId(name);",
            "  return rpGrades.find((g) => normalizeGradeId(g.name) === key);",
            "}",
            "",
            "export const rpGradesByBranch = Object.groupBy(rpGrades, (g) => g.branch);",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    from discord_catalog_gen import (
        build_discord_catalog,
        render_discord_master_ts,
        render_discord_md,
    )

    excel = resolve_excel_path(sys.argv)
    names, grades = from_excel(excel)
    names = dedupe_names(names)
    source = excel.name if excel.exists() else f"Google Sheets {SHEETS_ID}"
    catalog = render_catalog_ts(names, source)
    web = render_web_ts(grades, source)
    for out in (OUT_BOT, OUT_API):
        out.write_text(catalog, encoding="utf-8")
        print(f"Wrote {len(names)} grades -> {out}")
    OUT_WEB.write_text(web, encoding="utf-8")
    print(f"Wrote {len(grades)} grade metadata -> {OUT_WEB}")

    foundation, sections, display_map = build_discord_catalog(names, grades, normalize)
    discord_ts = render_discord_master_ts(foundation, sections, display_map, source)
    discord_md = render_discord_md(sections, source)
    OUT_DISCORD_GEN.write_text(discord_ts, encoding="utf-8")
    OUT_DISCORD_MD.write_text(discord_md, encoding="utf-8")
    print(f"Wrote {len(foundation)} discord roles -> {OUT_DISCORD_GEN}")
    print(f"Wrote guide -> {OUT_DISCORD_MD}")
    print("\nSync terminee : Excel -> Site + Bot + API + Guide Discord")


if __name__ == "__main__":
    main()
